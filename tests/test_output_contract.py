"""Cross-cutting regression: asserts the output *contract* at the outermost
artifacts an analyst or an API consumer actually sees -- the parsed CSV row
and the API's `/scores` JSON -- not the internal objects that build them.

Four of the nine findings in docs/2026-09-02-codebase-evaluation.md were
invisible to a 167-test suite because the suite tested the function that
implements a guarantee (e.g. matched_field being populated on ScreeningHit)
rather than the boundary that's supposed to deliver it (e.g. matched_field
actually reaching the CSV a reviewer opens). This file exists to prevent
that class of gap recurring -- for matched_field/producer (Finding 5),
attribution/license/caveat (Finding 6), and the "never confirmed" language
discipline (Section 10) all at once, not one regression test per instance.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from entity_screening.api.main import app
from entity_screening.common import storage
from entity_screening.common.attribution import OPENALEX_PRECISION_CAVEAT, attribution_for
from entity_screening.common.schema import (
    CaseState,
    ForeignControlFlag,
    MatchStatus,
    ScreeningHit,
    WorksheetActionKind,
)
from entity_screening.common.schema import _FORBIDDEN_OBSERVATION_FIELD_TOKENS

FIXTURES_DIR = Path(__file__).parent / "fixtures"
NSF_FILE = str(FIXTURES_DIR / "sample_nsf_awards.json")
OPENSANCTIONS_FILE = str(FIXTURES_DIR / "sample_opensanctions_targets.csv")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("ENTITY_SCREENING_DB_PATH", str(tmp_path / "test.duckdb"))
    monkeypatch.setenv("ENTITY_SCREENING_RUNS_DIR", str(tmp_path / "runs"))
    return TestClient(app)


def _create_run(client) -> str:
    response = client.post("/runs", json={"opensanctions_file": OPENSANCTIONS_FILE, "nsf_file": NSF_FILE})
    assert response.status_code == 200, response.text
    return response.json()["run_id"]


def _hit(entity_id: str, list_name: str, producer: str, confidence: float = 0.95) -> ScreeningHit:
    """Builds evidence["source_attribution"] the same way the real producers
    do (screen.py/section_117.py/cross_check.py), rather than hand-typing
    fake attribution text, so this test exercises the real contract."""
    source_attribution = attribution_for(list_name)
    if producer == "bibliometric":
        source_attribution = {**source_attribution, "caveat": OPENALEX_PRECISION_CAVEAT}
    return ScreeningHit(
        entity_id=entity_id,
        list_name=list_name,
        matched_variant="Fixture Matched Variant",
        matched_field="name_variants",
        confidence=confidence,
        evidence={"entry_id": "fixture-entry-1", "source_attribution": source_attribution},
        status=MatchStatus.CANDIDATE_MATCH,
        producer=producer,
    )


def test_output_contract_at_the_csv_and_api_boundaries(client, tmp_path):
    run_id = _create_run(client)
    entity_id = client.get(f"/runs/{run_id}/scores").json()[0]["entity_id"]

    hits = [
        _hit(entity_id, "opensanctions_consolidated", "direct_name"),
        _hit(entity_id, "section_117_foreign_funding_disclosure", "section_117"),
        _hit(entity_id, "dod_section_1260h", "bibliometric"),
    ]
    flag = ForeignControlFlag(
        entity_id=entity_id,
        entity_lei="LEI-FIXTURE-1",
        entity_jurisdiction="US",
        ultimate_parent_lei="LEI-FIXTURE-2",
        ultimate_parent_name="Fixture Foreign Parent",
        ultimate_parent_jurisdiction="DE",
        relationship_path=("LEI-FIXTURE-1", "LEI-FIXTURE-2"),
        match_confidence=0.9,
        evidence={"lei_match_basis": "normalized_exact", "relationship_path": ["LEI-FIXTURE-1", "LEI-FIXTURE-2"], "truncated": False},
        status=MatchStatus.CANDIDATE_MATCH,
    )

    conn = storage.connect(tmp_path / "test.duckdb")
    storage.insert_screening_hits(conn, hits, run_id)
    storage.insert_ownership_flags(conn, [flag], run_id)
    conn.close()

    # ---- API boundary: GET /runs/{id}/scores ----
    scores = client.get(f"/runs/{run_id}/scores").json()
    entity = next(s for s in scores if s["entity_id"] == entity_id)
    assert entity["status"] == "candidate_match"
    assert len(entity["screening_hits"]) == 3
    _assert_hits_carry_the_full_contract(entity["screening_hits"])

    # ---- CSV boundary: GET /runs/{id}/export.csv ----
    csv_response = client.get(f"/runs/{run_id}/export.csv")
    assert csv_response.status_code == 200
    reader = csv.DictReader(io.StringIO(csv_response.text))
    row = next(r for r in reader if r["entity_id"] == entity_id)
    assert row["status"] == "candidate_match"
    csv_hits = json.loads(row["screening_hits"])
    assert len(csv_hits) == 3
    _assert_hits_carry_the_full_contract(csv_hits)

    # ---- Language discipline (Section 10): never "confirmed" anywhere ----
    assert "confirmed" not in csv_response.text.lower()
    assert "confirmed" not in json.dumps(scores).lower()

    # ---- status is candidate_match iff hits or ownership flags exist ----
    for s in scores:
        has_evidence = bool(s["screening_hits"]) or bool(s["ownership_flags"])
        assert s["status"] == ("candidate_match" if has_evidence else "no_hit")


def _walk_keys(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield k
            yield from _walk_keys(v)
    elif isinstance(obj, list):
        for item in obj:
            yield from _walk_keys(item)


def test_investigative_file_export_contract(tmp_path):
    """The investigative file is a new output boundary (use-case-01). The
    same guarantees the batch CSV carries must hold here: attribution +
    licence on every evidence payload (Section 10), no evaluative field
    anywhere in the Finding graph as serialized (Section 4), and the
    language discipline. Plus: classified subject fields redacted by
    default (Section 9)."""
    from entity_screening.case import demo, export, service, store

    db_path = tmp_path / "case.duckdb"
    conn = storage.connect(db_path)
    demo.build_demo_case(conn)
    conn.close()

    from entity_screening.pipeline import reconcile_case

    reconcile_case(
        "demo",
        db_path=db_path,
        runs_dir=tmp_path / "runs",
        works_fixture=demo.load_demo_works_fixture(),
        gleif_lei_file=demo.DEMO_GLEIF_LEI_FILE,
        gleif_relationships_file=demo.DEMO_GLEIF_RELATIONSHIPS_FILE,
    )

    conn = storage.connect(db_path)
    view = service.worksheet(conn, "demo")
    # Discrepancy rows: dismiss all but one; certification-required on the last.
    for row in view.rows[:-1]:
        service.record_action(
            conn, "demo", row.finding.finding_id, WorksheetActionKind.DISMISS,
            "record_error_or_misattribution", "stale", "analyst.a",
        )
    last = view.rows[-1]
    service.record_action(
        conn, "demo", last.finding.finding_id, WorksheetActionKind.CERTIFICATION_REQUIRED,
        "possible_nondisclosure_for_certification", "undisclosed affiliation", "analyst.a",
    )
    service.record_certification(
        conn, "demo", last.finding.finding_id,
        "Undisclosed affiliation to a foreign institution.",
        "Collaboration was publicly documented and disclosed elsewhere in the packet.",
        "Dr. Department Head",
    )
    # Concern-tie rows must be actioned too, or the case cannot close.
    for tie_row in view.tie_rows:
        service.record_tie_action(
            conn, "demo", tie_row.tie.tie_id, WorksheetActionKind.ESCALATE,
            "needs_counterintelligence_referral", "ultimate parent on 1260H", "analyst.a",
        )
    service.transition(conn, "demo", CaseState.ADJUDICATION)
    service.record_adjudication(conn, "demo", "One item routed for certification; a tie escalated.", "Proceed with certification on file.", "analyst.a")

    out_path, manifest = export.export_investigative_file(
        conn, "demo", fmt="json", runs_dir=tmp_path / "runs"
    )
    unredacted_path, _ = export.export_investigative_file(
        conn, "demo", fmt="json", redact=False, runs_dir=tmp_path / "runs"
    )
    xlsx_path, _ = export.export_investigative_file(
        conn, "demo", fmt="xlsx", runs_dir=tmp_path / "runs"
    )
    conn.close()

    payload = json.loads(out_path.read_text())

    # --- the synthetic marker survives to the export, in both formats ---
    # (the file is the artifact that leaves the system; a downloaded JSON/XLSX
    # naming a real 1260H company in a fabricated ownership chain must say so)
    assert payload["provenance"]["synthetic"] is True
    notice = payload["provenance"]["notice"].lower()
    assert "synthetic" in notice and "fabricated" in notice
    assert "not a real finding" in notice
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path)
    assert "READ ME -- provenance" in wb.sheetnames
    provenance_text = " ".join(
        str(c.value) for row in wb["READ ME -- provenance"].iter_rows() for c in row
    ).lower()
    assert "synthetic" in provenance_text and "not a real finding" in provenance_text
    # The "Concern ties" sheet sits before "Findings" -- higher-stakes rows first.
    assert wb.sheetnames.index("Concern ties") < wb.sheetnames.index("Findings")
    # Header rows render even when a sheet has zero data rows.
    for name in ("Adjudications", "Certifications", "Concern ties", "Tie actions"):
        assert wb[name].max_row >= 1

    # --- every finding row carries the discrepancy contract (no evidence now) ---
    assert payload["findings"]
    for finding in payload["findings"]:
        assert finding["discovered"]["institution_name"]
        assert finding["declaration_search"]
        assert finding["factual_basis"]
        assert "concern_list_evidence" not in finding
        assert "ownership_evidence" not in finding

    # --- every concern-tie evidence payload carries attribution + licence (Section 10) ---
    assert payload["concern_ties"]
    for tie in payload["concern_ties"]:
        for ev in list(tie["concern_list_evidence"]) + list(tie["ownership_evidence"]):
            attribution = ev["evidence"]["source_attribution"]
            assert attribution["attribution"], "Section 10: attribution must reach the investigative file"
            assert attribution["license"], "Section 10: licence must reach the investigative file"

    # --- no evaluative field anywhere in the serialized observation graphs (Section 4) ---
    observation_keys = set(_walk_keys(payload["findings"])) | set(
        _walk_keys(
            [
                {k: v for k, v in t.items() if k != "concern_list_evidence"}
                for t in payload["concern_ties"]
            ]
        )
    )
    for key in observation_keys:
        for token in _FORBIDDEN_OBSERVATION_FIELD_TOKENS:
            assert token not in key.lower(), f"evaluative key {key!r} in the exported observations"

    # --- language discipline: system-generated text only (reason_note / notes are human) ---
    system_text = json.dumps(
        {
            k: v
            for k, v in payload.items()
            if k not in ("worksheet", "tie_actions", "adjudications", "certifications")
        }
    ).lower()
    assert "confirmed" not in system_text
    assert "risk score" not in system_text

    # --- redaction (Section 9) ---
    assert payload["subject"]["classified_fields"] == {
        "_redacted": True,
        "_reason": "field-level sensitive (use-case-01 Section 9)",
    }
    assert "SYNTH-000000" not in out_path.read_text()  # passport number, redacted
    assert "SYNTH-000000" in unredacted_path.read_text()  # present only when asked

    assert manifest.redaction_profile == "default"
    assert manifest.finding_count == len(payload["findings"])
    assert xlsx_path.exists()


def _assert_hits_carry_the_full_contract(hits: list[dict]) -> None:
    seen_producers = set()
    for hit in hits:
        assert hit["matched_field"], "matched_field must reach every hit output (Finding 5)"
        assert hit["producer"] in ("direct_name", "section_117", "bibliometric")
        assert hit["list_name"]
        assert hit["matched_variant"]
        assert isinstance(hit["confidence"], float)
        assert hit["status"]
        seen_producers.add(hit["producer"])

        source_attribution = hit["evidence"]["source_attribution"]
        assert source_attribution["attribution"], "attribution must reach every hit's evidence (Finding 6)"
        assert source_attribution["license"], "license must reach every hit's evidence (Section 10)"
        if hit["producer"] == "bibliometric":
            assert source_attribution.get("caveat"), (
                "every bibliometric hit's evidence must carry the OpenAlex precision caveat"
            )
        else:
            assert "caveat" not in source_attribution, (
                "only a bibliometric hit inherits OpenAlex's own uncertainty -- a direct_name "
                "or section_117 hit's attribution should not carry a caveat it didn't earn"
            )
    assert seen_producers == {"direct_name", "section_117", "bibliometric"}
