"""M8: the investigative-file export's provenance notice used to be the
demo's own GLEIF/SYNTH-LEI-specific text (case/export.py:PROVENANCE_NOTICE),
stamped unconditionally on every case's export -- including a non-demo case
created via POST /cases, which has nothing to do with the demo's bundled
GLEIF fixture.

Also M9: the "Declared affiliations" XLSX sheet is the one sheet in
_write_xlsx that skipped the shared `sheet()` helper, so it renders with no
header row at all when a declaration has zero affiliations.
"""
from __future__ import annotations

from datetime import date

import openpyxl

from entity_screening.case import demo, export as case_export, store as case_store
from entity_screening.common import storage
from entity_screening.common.schema import (
    Case,
    CaseKind,
    CaseState,
    CoverageBasis,
    Declaration,
    Subject,
)


def _non_demo_subject_declaration_case(case_id: str):
    subject = Subject(
        subject_id=f"{case_id}-subject",
        display_name="Test Subject",
        coverage_basis=CoverageBasis.FOREIGN_ADVERSARY_TIE,
        synthetic=True,
        classified_fields={},
    )
    declaration = Declaration(
        declaration_id=f"{case_id}-declaration",
        subject_id=subject.subject_id,
        synthetic=True,
        sources=(),
        affiliations=(),
    )
    case = Case(
        case_id=case_id,
        subject_id=subject.subject_id,
        declaration_id=declaration.declaration_id,
        trigger="hire",
        access_scope="data",
        coverage_basis=CoverageBasis.FOREIGN_ADVERSARY_TIE,
        synthetic=True,
        case_kind=CaseKind.HB127_RESEARCHER_SCREENING,
        state=CaseState.INTAKE,
        statutory_deadline=date(2026, 10, 1),
    )
    return subject, declaration, case


def test_demo_case_export_keeps_the_exact_gleif_specific_notice(tmp_path):
    conn = storage.connect(tmp_path / "case.duckdb")
    demo.build_demo_case(conn)
    payload = case_export.build_investigative_file(conn, demo.DEMO_CASE_ID)
    conn.close()

    assert payload["provenance"]["notice"] == case_export.DEMO_PROVENANCE_NOTICE
    assert "SYNTH" in payload["provenance"]["notice"]
    assert "gleif.NOTICE.md" in payload["provenance"]["notice"]


def test_non_demo_case_export_gets_the_generic_notice_not_the_demo_text(tmp_path):
    conn = storage.connect(tmp_path / "case.duckdb")
    subject, declaration, case = _non_demo_subject_declaration_case("c-non-demo")
    case_store.save_subject(conn, subject)
    case_store.save_declaration(conn, declaration)
    case_store.save_case(conn, case)

    payload = case_export.build_investigative_file(conn, "c-non-demo")
    conn.close()

    assert payload["provenance"]["notice"] == case_export.GENERIC_PROVENANCE_NOTICE
    assert payload["provenance"]["notice"] != case_export.DEMO_PROVENANCE_NOTICE
    # The generic notice must not carry demo-fixture-specific claims that
    # aren't true for this case.
    assert "LEIs prefixed" not in payload["provenance"]["notice"]
    assert "gleif.NOTICE.md" not in payload["provenance"]["notice"]


def test_declared_affiliations_sheet_renders_its_header_with_zero_rows(tmp_path):
    conn = storage.connect(tmp_path / "case.duckdb")
    subject, declaration, case = _non_demo_subject_declaration_case("c-empty-decl")
    case_store.save_subject(conn, subject)
    case_store.save_declaration(conn, declaration)  # zero affiliations
    case_store.save_case(conn, case)

    out_path, _ = case_export.export_investigative_file(
        conn, "c-empty-decl", fmt="xlsx", runs_dir=tmp_path / "runs"
    )
    conn.close()

    wb = openpyxl.load_workbook(out_path)
    sheet = wb["Declared affiliations"]
    assert sheet.max_row >= 1  # the header row itself, even with zero data rows
    header = [c.value for c in next(sheet.iter_rows(max_row=1))]
    assert "institution_name" in header
